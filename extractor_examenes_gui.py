#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extractor de exámenes ocupacionales (PDF -> Excel) con interfaz gráfica.

- Elige la carpeta de origen con los PDF.
- Marca si deseas incluir las subcarpetas.
- Elige la carpeta de destino y el nombre del Excel.
- Presiona "Procesar".

Formatos reconocidos: CMT (Centro Médico del Trabajador) y WORKMED.
Cualquier otro PDF queda marcado como "FORMATO NO RECONOCIDO" en el Excel.

Funciona 100 % sin conexión a internet. No usa inteligencia artificial.

Requisitos (instalar una sola vez):
    pip install pdfplumber openpyxl

Ejecución:
    python extractor_examenes_gui.py
"""

import queue
import re
import sys
import threading
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as e:
    root = tk.Tk(); root.withdraw()
    messagebox.showerror(
        "Falta una librería",
        f"No se encontró la librería: {e.name}\n\n"
        "Instálala con:\n    pip install pdfplumber openpyxl")
    sys.exit(1)


# ===========================================================================
# LÓGICA DE EXTRACCIÓN
# ===========================================================================
COLUMNAS = [
    ("archivo",    "Archivo PDF"),
    ("centro",     "Centro médico"),
    ("nombre",     "Nombre completo"),
    ("rut",        "RUT"),
    ("cargo",      "Cargo"),
    ("codigo",     "Código de verificación (QR)"),
    ("reserva",    "N° de reserva"),
    ("examen",     "Examen"),
    ("fecha_eval", "Fecha de evaluación"),
    ("fecha_venc", "Vigencia hasta"),
    ("resultado",  "Resultado"),
    ("sucursal",   "Sucursal / lugar"),
    ("estado",     "Estado extracción"),
]

MESES = {"enero": "01", "febrero": "02", "marzo": "03", "abril": "04",
         "mayo": "05", "junio": "06", "julio": "07", "agosto": "08",
         "septiembre": "09", "setiembre": "09", "octubre": "10",
         "noviembre": "11", "diciembre": "12"}


def _buscar(patron, texto, flags=re.IGNORECASE):
    m = re.search(patron, texto, flags)
    return " ".join(m.group(1).split()) if m else ""


def _fecha_larga_a_corta(s):
    """'28 de Septiembre de 2027' -> '28-09-2027'"""
    m = re.match(r"(\d{1,2})\s+de\s+(\w+)\s+de\s+(\d{4})", s, re.IGNORECASE)
    if not m:
        return s
    mes = MESES.get(m.group(2).lower(), "??")
    return f"{int(m.group(1)):02d}-{mes}-{m.group(3)}"


# ---------------------------------------------------------------------------
# Parsers por formato (reciben el texto de UNA página)
# ---------------------------------------------------------------------------
def parsear_cmt(texto):
    """Página de portada de un examen CMT (contiene el código QR)."""
    return {
        "centro":     "CMT",
        "nombre":     _buscar(r"Nombre:\s*(.+?)\s+RUT:", texto),
        "rut":        _buscar(r"Nombre:.+?RUT:\s*([\d\.]+-[\dkK])", texto, re.I | re.S),
        "cargo":      _buscar(r"Cargo:\s*(.+?)\s*(?:\n|$)", texto),
        "codigo":     _buscar(r"c[oó]digo:\s*([A-Za-z0-9]+)", texto),
        "reserva":    _buscar(r"N[°º]\s*de\s*reserva:\s*([A-Za-z0-9]+)", texto),
        "examen":     _buscar(r"BATER[IÍ]A\s+DE\s+RIESGO:\s*(.+?)\s+Fecha de evaluaci", texto),
        "fecha_eval": _buscar(r"Fecha de evaluaci[oó]n:\s*(\d{2}-\d{2}-\d{4})", texto),
        "fecha_venc": _buscar(r"Fecha de vencimiento:\s*(\d{2}-\d{2}-\d{4})", texto),
        "resultado":  _buscar(r"RESULTADO DE EVALUACI[OÓ]N:\s*([A-ZÁÉÍÓÚÑ ]+?)\s+Fecha", texto),
        "sucursal":   _buscar(r"Sucursal de evaluaci[oó]n:\s*(.+?)\s*(?:\n|$)", texto),
    }


def parsear_workmed(texto):
    """Página de un examen WORKMED (evaluación laboral o screening A&D)."""
    examen = _buscar(r"Evaluaci[oó]n Laboral\s*\n\s*(.+?)\s*\n", texto)
    fecha = (_buscar(r"Fecha de Evaluaci[oó]n:\s*(\d{2}-\d{2}-\d{4})", texto)
             or _buscar(r"Fecha:\s*(\d{2}-\d{2}-\d{4})", texto))
    cargo = (_buscar(r"Cargo al que postula/ocupa\s*:?\s*(.+?)\s*(?:\n|$)", texto)
             or _buscar(r"\nCargo:\s*(.+?)\s*(?:\n|$)", texto))
    vig = _buscar(r"VIGENCIA HASTA:\s*(.+?)\s*(?:\n|$)", texto)
    resultado = _buscar(r"RESULTADO:\s*([A-ZÁÉÍÓÚÑ ]+?)\.\s", texto, flags=0)  # "APTO." / "NO APTO."
    if not resultado and "SCREENING" in examen.upper():
        # En el screening no hay campo RESULTADO; se infiere de las marcas.
        # ◉ en la columna "Resultado Negativo" para todas las drogas.
        no_neg = re.findall(r"\(\w{2,3}\)\s*○\s*◉", texto)
        resultado = "NO NEGATIVO" if no_neg else "NEGATIVO"
    return {
        "centro":     "WORKMED",
        "nombre":     _buscar(r"Nombre:\s*(.+?)\s+Nombre:", texto),
        "rut":        _buscar(r"\nRUT:\s*([\d\.]+-[\dkK])", texto),
        "cargo":      cargo,
        "codigo":     _buscar(r"Folio ID:\s*([\w-]+)", texto),
        "reserva":    "",                       # WORKMED no emite N° de reserva
        "examen":     examen,
        "fecha_eval": fecha,
        "fecha_venc": _fecha_larga_a_corta(vig),
        "resultado":  resultado,
        "sucursal":   _buscar(r"Lugar de Evaluaci[oó]n:\s*(.+?)\s+(?:Fecha de Emisi|Hora toma)", texto),
    }


def detectar_formato(texto):
    if "con el siguiente código" in texto or "con el siguiente codigo" in texto:
        return "cmt"
    if "workmed" in texto.lower() and "Folio ID" in texto:
        return "workmed"
    return None


# ---------------------------------------------------------------------------
# Procesamiento de un PDF
# ---------------------------------------------------------------------------
def extraer_pdf(ruta: Path) -> list[dict]:
    filas = []
    try:
        with pdfplumber.open(ruta) as pdf:
            paginas = [(p.extract_text() or "") for p in pdf.pages]
    except Exception as e:
        return [{**{c: "" for c, _ in COLUMNAS}, "archivo": ruta.name,
                 "estado": f"ERROR al abrir: {e}"}]

    if not any(t.strip() for t in paginas):
        return [{**{c: "" for c, _ in COLUMNAS}, "archivo": ruta.name,
                 "estado": "SIN TEXTO (PDF escaneado, requiere OCR)"}]

    vistos = set()
    for texto in paginas:
        fmt = detectar_formato(texto)
        if fmt == "cmt":
            datos = parsear_cmt(texto)
        elif fmt == "workmed":
            datos = parsear_workmed(texto)
        else:
            continue  # páginas de resultados detallados, sin datos de cabecera

        clave = (datos["codigo"], datos["examen"])
        if not datos["codigo"] or clave in vistos:
            continue
        vistos.add(clave)

        fila = {c: "" for c, _ in COLUMNAS}
        fila.update(datos)
        fila["archivo"] = ruta.name
        fila["nombre"] = fila["nombre"].upper()
        obligatorios = ["nombre", "rut", "cargo", "codigo", "examen", "fecha_eval"]
        faltan = [c for c in obligatorios if not fila[c]]
        fila["estado"] = "OK" if not faltan else "REVISAR: falta " + ", ".join(faltan)
        filas.append(fila)

    if not filas:
        filas.append({**{c: "" for c, _ in COLUMNAS}, "archivo": ruta.name,
                      "estado": "FORMATO NO RECONOCIDO"})
    return filas


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def escribir_excel(filas, salida: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exámenes"

    fuente = Font(name="Arial", size=10)
    fuente_cab = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    relleno_cab = PatternFill("solid", fgColor="1F4E78")
    relleno_rev = PatternFill("solid", fgColor="FFF2CC")

    ws.append([t for _, t in COLUMNAS])
    for c in ws[1]:
        c.font, c.fill = fuente_cab, relleno_cab
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for fila in filas:
        ws.append([fila[c] for c, _ in COLUMNAS])
        for c in ws[ws.max_row]:
            c.font = fuente
            if fila["estado"] != "OK":
                c.fill = relleno_rev

    for i, (clave, titulo) in enumerate(COLUMNAS, start=1):
        largo = max([len(titulo)] + [len(str(f[clave])) for f in filas])
        ws.column_dimensions[get_column_letter(i)].width = min(largo + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(salida)



def procesar_carpeta(origen: Path, recursivo: bool, salida: Path, log, progreso):
    """Ejecuta la extracción completa. `log(msg)` y `progreso(i, n)` son
    callbacks para informar a la interfaz."""
    pdfs = sorted(origen.rglob("*.pdf") if recursivo else origen.glob("*.pdf"))
    if not pdfs:
        raise FileNotFoundError(f"No se encontraron archivos PDF en {origen}")

    filas = []
    for i, pdf in enumerate(pdfs, start=1):
        nuevas = extraer_pdf(pdf)
        # Guardar ruta relativa para ubicar el archivo si hay subcarpetas
        for f in nuevas:
            f["archivo"] = str(pdf.relative_to(origen))
        filas.extend(nuevas)
        ok = sum(1 for f in nuevas if f["estado"] == "OK")
        log(f"[{i}/{len(pdfs)}] {pdf.name}: {len(nuevas)} examen(es), {ok} OK")
        progreso(i, len(pdfs))

    escribir_excel(filas, salida)
    total_ok = sum(1 for f in filas if f["estado"] == "OK")
    no_rec = sum(1 for f in filas if f["estado"] == "FORMATO NO RECONOCIDO")
    return len(pdfs), len(filas), total_ok, no_rec


# ===========================================================================
# INTERFAZ GRÁFICA
# ===========================================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Extractor de exámenes ocupacionales - PDF a Excel")
        self.resizable(True, False)
        self.cola = queue.Queue()
        self._construir()
        self.after(100, self._vaciar_cola)

    def _construir(self):
        pad = {"padx": 8, "pady": 4}
        marco = ttk.Frame(self, padding=10)
        marco.pack(fill="both", expand=True)
        marco.columnconfigure(1, weight=1)

        # Origen
        ttk.Label(marco, text="Carpeta de origen (PDF):").grid(row=0, column=0, sticky="w", **pad)
        self.var_origen = tk.StringVar()
        ttk.Entry(marco, textvariable=self.var_origen).grid(row=0, column=1, sticky="ew", **pad)
        ttk.Button(marco, text="Buscar...", command=self._elegir_origen).grid(row=0, column=2, **pad)

        # Subcarpetas
        self.var_recursivo = tk.BooleanVar(value=True)
        ttk.Checkbutton(marco, text="Incluir archivos de las subcarpetas",
                        variable=self.var_recursivo).grid(row=1, column=1, sticky="w", **pad)

        # Destino
        ttk.Label(marco, text="Carpeta de destino (Excel):").grid(row=2, column=0, sticky="w", **pad)
        self.var_destino = tk.StringVar()
        ttk.Entry(marco, textvariable=self.var_destino).grid(row=2, column=1, sticky="ew", **pad)
        ttk.Button(marco, text="Buscar...", command=self._elegir_destino).grid(row=2, column=2, **pad)

        # Nombre archivo
        ttk.Label(marco, text="Nombre del archivo Excel:").grid(row=3, column=0, sticky="w", **pad)
        self.var_nombre = tk.StringVar(value="examenes_ocupacionales.xlsx")
        ttk.Entry(marco, textvariable=self.var_nombre).grid(row=3, column=1, sticky="ew", **pad)

        # Botón procesar
        self.btn = ttk.Button(marco, text="Procesar", command=self._procesar)
        self.btn.grid(row=4, column=1, sticky="e", **pad)

        # Barra de progreso
        self.barra = ttk.Progressbar(marco, mode="determinate")
        self.barra.grid(row=5, column=0, columnspan=3, sticky="ew", **pad)

        # Registro
        ttk.Label(marco, text="Registro:").grid(row=6, column=0, sticky="nw", **pad)
        self.txt = tk.Text(marco, height=14, width=90, state="disabled", wrap="none")
        self.txt.grid(row=6, column=1, columnspan=2, sticky="nsew", **pad)
        sb = ttk.Scrollbar(marco, orient="vertical", command=self.txt.yview)
        sb.grid(row=6, column=3, sticky="ns")
        self.txt.configure(yscrollcommand=sb.set)

        ttk.Label(marco, foreground="gray",
                  text="Formatos reconocidos: CMT y WORKMED. Otros PDF se marcan como "
                       "FORMATO NO RECONOCIDO. Funciona sin internet.").grid(
            row=7, column=0, columnspan=3, sticky="w", **pad)

    # --- selección de carpetas -------------------------------------------
    def _elegir_origen(self):
        d = filedialog.askdirectory(title="Selecciona la carpeta con los PDF")
        if d:
            self.var_origen.set(d)
            if not self.var_destino.get():
                self.var_destino.set(d)

    def _elegir_destino(self):
        d = filedialog.askdirectory(title="Selecciona la carpeta donde guardar el Excel")
        if d:
            self.var_destino.set(d)

    # --- registro y progreso (seguros entre hilos) ------------------------
    def _log(self, msg):
        self.cola.put(("log", msg))

    def _progreso(self, i, n):
        self.cola.put(("prog", (i, n)))

    def _vaciar_cola(self):
        try:
            while True:
                tipo, dato = self.cola.get_nowait()
                if tipo == "log":
                    self.txt.configure(state="normal")
                    self.txt.insert("end", dato + "\n")
                    self.txt.see("end")
                    self.txt.configure(state="disabled")
                elif tipo == "prog":
                    i, n = dato
                    self.barra["maximum"] = n
                    self.barra["value"] = i
                elif tipo == "fin":
                    self.btn.configure(state="normal")
                    n_pdf, n_filas, n_ok, n_norec = dato
                    messagebox.showinfo(
                        "Proceso terminado",
                        f"PDF procesados: {n_pdf}\n"
                        f"Exámenes extraídos: {n_filas}\n"
                        f"Sin observaciones: {n_ok}\n"
                        f"Formato no reconocido: {n_norec}\n\n"
                        f"Excel guardado en:\n{self.ruta_salida}")
                elif tipo == "error":
                    self.btn.configure(state="normal")
                    messagebox.showerror("Error", dato)
        except queue.Empty:
            pass
        self.after(100, self._vaciar_cola)

    # --- ejecución ---------------------------------------------------------
    def _procesar(self):
        origen = Path(self.var_origen.get().strip())
        destino = Path(self.var_destino.get().strip())
        nombre = self.var_nombre.get().strip() or "examenes_ocupacionales.xlsx"
        if not nombre.lower().endswith(".xlsx"):
            nombre += ".xlsx"

        if not origen.is_dir():
            messagebox.showwarning("Falta información", "Selecciona una carpeta de origen válida.")
            return
        if not destino.is_dir():
            messagebox.showwarning("Falta información", "Selecciona una carpeta de destino válida.")
            return

        self.ruta_salida = destino / nombre
        if self.ruta_salida.exists():
            if not messagebox.askyesno("Archivo existente",
                                       f"{nombre} ya existe en la carpeta de destino.\n¿Deseas reemplazarlo?"):
                return

        self.txt.configure(state="normal"); self.txt.delete("1.0", "end"); self.txt.configure(state="disabled")
        self.barra["value"] = 0
        self.btn.configure(state="disabled")

        def tarea():
            try:
                res = procesar_carpeta(origen, self.var_recursivo.get(),
                                       self.ruta_salida, self._log, self._progreso)
                self.cola.put(("fin", res))
            except PermissionError:
                self.cola.put(("error", "No se pudo guardar el Excel. "
                                        "Si está abierto en Excel, ciérralo e intenta de nuevo."))
            except Exception as e:
                self.cola.put(("error", str(e)))

        threading.Thread(target=tarea, daemon=True).start()


if __name__ == "__main__":
    App().mainloop()
