#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extrae los datos principales de exámenes ocupacionales en PDF y los vuelca
en un Excel, una fila por examen.

Formatos reconocidos:
  - CMT (Centro Médico del Trabajador): código de verificación del QR
    ("con el siguiente código: XXXXX") + N° de reserva.
  - WORKMED: el código de verificación es el "Folio ID" del pie de página.

Un PDF puede contener varios exámenes (altura, ruido, sílice, alcohol y
drogas...) y cada examen puede ocupar varias páginas; el script detecta
cada examen y elimina duplicados por código.

Uso:
    python extraer_examenes.py <carpeta_con_pdf> [salida.xlsx]

Requisitos:  pip install pdfplumber openpyxl
"""

import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNAS = [
    ("archivo",    "Archivo PDF"),
    ("centro",     "Centro médico"),
    ("nombre",     "Nombre completo"),
    ("rut",        "RUT"),
    ("cargo",      "Cargo"),
    ("codigo",     "Código de verificación (QR)"),
    ("reserva",    "N° de reserva"),
    ("examen",     "Examen"),
    ("fecha_eval", "Fecha de evaluación"),
    ("fecha_venc", "Vigencia hasta"),
    ("resultado",  "Resultado"),
    ("sucursal",   "Sucursal / lugar"),
    ("estado",     "Estado extracción"),
]

MESES = {"enero": "01", "febrero": "02", "marzo": "03", "abril": "04",
         "mayo": "05", "junio": "06", "julio": "07", "agosto": "08",
         "septiembre": "09", "setiembre": "09", "octubre": "10",
         "noviembre": "11", "diciembre": "12"}


def _buscar(patron, texto, flags=re.IGNORECASE):
    m = re.search(patron, texto, flags)
    return " ".join(m.group(1).split()) if m else ""


def _fecha_larga_a_corta(s):
    """'28 de Septiembre de 2027' -> '28-09-2027'"""
    m = re.match(r"(\d{1,2})\s+de\s+(\w+)\s+de\s+(\d{4})", s, re.IGNORECASE)
    if not m:
        return s
    mes = MESES.get(m.group(2).lower(), "??")
    return f"{int(m.group(1)):02d}-{mes}-{m.group(3)}"


# ---------------------------------------------------------------------------
# Parsers por formato (reciben el texto de UNA página)
# ---------------------------------------------------------------------------
def parsear_cmt(texto):
    """Página de portada de un examen CMT (contiene el código QR)."""
    return {
        "centro":     "CMT",
        "nombre":     _buscar(r"Nombre:\s*(.+?)\s+RUT:", texto),
        "rut":        _buscar(r"Nombre:.+?RUT:\s*([\d\.]+-[\dkK])", texto, re.I | re.S),
        "cargo":      _buscar(r"Cargo:\s*(.+?)\s*(?:\n|$)", texto),
        "codigo":     _buscar(r"c[oó]digo:\s*([A-Za-z0-9]+)", texto),
        "reserva":    _buscar(r"N[°º]\s*de\s*reserva:\s*([A-Za-z0-9]+)", texto),
        "examen":     _buscar(r"BATER[IÍ]A\s+DE\s+RIESGO:\s*(.+?)\s+Fecha de evaluaci", texto),
        "fecha_eval": _buscar(r"Fecha de evaluaci[oó]n:\s*(\d{2}-\d{2}-\d{4})", texto),
        "fecha_venc": _buscar(r"Fecha de vencimiento:\s*(\d{2}-\d{2}-\d{4})", texto),
        "resultado":  _buscar(r"RESULTADO DE EVALUACI[OÓ]N:\s*([A-ZÁÉÍÓÚÑ ]+?)\s+Fecha", texto),
        "sucursal":   _buscar(r"Sucursal de evaluaci[oó]n:\s*(.+?)\s*(?:\n|$)", texto),
    }


def parsear_workmed(texto):
    """Página de un examen WORKMED (evaluación laboral o screening A&D)."""
    examen = _buscar(r"Evaluaci[oó]n Laboral\s*\n\s*(.+?)\s*\n", texto)
    fecha = (_buscar(r"Fecha de Evaluaci[oó]n:\s*(\d{2}-\d{2}-\d{4})", texto)
             or _buscar(r"Fecha:\s*(\d{2}-\d{2}-\d{4})", texto))
    cargo = (_buscar(r"Cargo al que postula/ocupa\s*:?\s*(.+?)\s*(?:\n|$)", texto)
             or _buscar(r"\nCargo:\s*(.+?)\s*(?:\n|$)", texto))
    vig = _buscar(r"VIGENCIA HASTA:\s*(.+?)\s*(?:\n|$)", texto)
    resultado = _buscar(r"RESULTADO:\s*([A-ZÁÉÍÓÚÑ ]+?)\.\s", texto, flags=0)  # "APTO." / "NO APTO."
    if not resultado and "SCREENING" in examen.upper():
        # En el screening no hay campo RESULTADO; se infiere de las marcas.
        # ◉ en la columna "Resultado Negativo" para todas las drogas.
        no_neg = re.findall(r"\(\w{2,3}\)\s*○\s*◉", texto)
        resultado = "NO NEGATIVO" if no_neg else "NEGATIVO"
    return {
        "centro":     "WORKMED",
        "nombre":     _buscar(r"Nombre:\s*(.+?)\s+Nombre:", texto),
        "rut":        _buscar(r"\nRUT:\s*([\d\.]+-[\dkK])", texto),
        "cargo":      cargo,
        "codigo":     _buscar(r"Folio ID:\s*([\w-]+)", texto),
        "reserva":    "",                       # WORKMED no emite N° de reserva
        "examen":     examen,
        "fecha_eval": fecha,
        "fecha_venc": _fecha_larga_a_corta(vig),
        "resultado":  resultado,
        "sucursal":   _buscar(r"Lugar de Evaluaci[oó]n:\s*(.+?)\s+(?:Fecha de Emisi|Hora toma)", texto),
    }


def detectar_formato(texto):
    if "con el siguiente código" in texto or "con el siguiente codigo" in texto:
        return "cmt"
    if "workmed" in texto.lower() and "Folio ID" in texto:
        return "workmed"
    return None


# ---------------------------------------------------------------------------
# Procesamiento de un PDF
# ---------------------------------------------------------------------------
def extraer_pdf(ruta: Path) -> list[dict]:
    filas = []
    try:
        with pdfplumber.open(ruta) as pdf:
            paginas = [(p.extract_text() or "") for p in pdf.pages]
    except Exception as e:
        return [{**{c: "" for c, _ in COLUMNAS}, "archivo": ruta.name,
                 "estado": f"ERROR al abrir: {e}"}]

    if not any(t.strip() for t in paginas):
        return [{**{c: "" for c, _ in COLUMNAS}, "archivo": ruta.name,
                 "estado": "SIN TEXTO (PDF escaneado, requiere OCR)"}]

    vistos = set()
    for texto in paginas:
        fmt = detectar_formato(texto)
        if fmt == "cmt":
            datos = parsear_cmt(texto)
        elif fmt == "workmed":
            datos = parsear_workmed(texto)
        else:
            continue  # páginas de resultados detallados, sin datos de cabecera

        clave = (datos["codigo"], datos["examen"])
        if not datos["codigo"] or clave in vistos:
            continue
        vistos.add(clave)

        fila = {c: "" for c, _ in COLUMNAS}
        fila.update(datos)
        fila["archivo"] = ruta.name
        fila["nombre"] = fila["nombre"].upper()
        obligatorios = ["nombre", "rut", "cargo", "codigo", "examen", "fecha_eval"]
        faltan = [c for c in obligatorios if not fila[c]]
        fila["estado"] = "OK" if not faltan else "REVISAR: falta " + ", ".join(faltan)
        filas.append(fila)

    if not filas:
        filas.append({**{c: "" for c, _ in COLUMNAS}, "archivo": ruta.name,
                      "estado": "FORMATO NO RECONOCIDO"})
    return filas


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def escribir_excel(filas, salida: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exámenes"

    fuente = Font(name="Arial", size=10)
    fuente_cab = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    relleno_cab = PatternFill("solid", fgColor="1F4E78")
    relleno_rev = PatternFill("solid", fgColor="FFF2CC")

    ws.append([t for _, t in COLUMNAS])
    for c in ws[1]:
        c.font, c.fill = fuente_cab, relleno_cab
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for fila in filas:
        ws.append([fila[c] for c, _ in COLUMNAS])
        for c in ws[ws.max_row]:
            c.font = fuente
            if fila["estado"] != "OK":
                c.fill = relleno_rev

    for i, (clave, titulo) in enumerate(COLUMNAS, start=1):
        largo = max([len(titulo)] + [len(str(f[clave])) for f in filas])
        ws.column_dimensions[get_column_letter(i)].width = min(largo + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(salida)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    carpeta = Path(sys.argv[1])
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else carpeta / "examenes.xlsx"

    pdfs = sorted(carpeta.rglob("*.pdf"))
    if not pdfs:
        print(f"No se encontraron PDF en {carpeta}")
        sys.exit(1)

    filas = []
    for i, pdf in enumerate(pdfs, start=1):
        nuevas = extraer_pdf(pdf)
        filas.extend(nuevas)
        ok = sum(1 for f in nuevas if f["estado"] == "OK")
        print(f"[{i}/{len(pdfs)}] {pdf.name}: {len(nuevas)} examen(es), {ok} OK")

    escribir_excel(filas, salida)
    ok = sum(1 for f in filas if f["estado"] == "OK")
    print(f"\nListo: {len(filas)} filas, {ok} sin observaciones. Excel: {salida}")


if __name__ == "__main__":
    main()
